import os
import re
import requests
from datetime import datetime
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Set your Discord Webhook URL here
DISCORD_WEBHOOK_URL = "https://discord.com/api/webhooks/1558535967504334954/d8e58aeo1uJIfuAurJ0OymPVxkdnn4v93Wg0r9_xvgF05-Hmi3Yg6t_RL26_A2Lt7KjL"  # Paste your Discord Webhook URL here


def send_discord_alert(device_name, source_file, failed_rules, webhook_url):
    """
    Dispatch a rich embed alert to Discord when security compliance checks fail.
    """
    if not webhook_url:
        print("    [!] Webhook URL not configured. Skipping alert.")
        return

    failures_formatted = "\n".join([f"• **{rule}**: {reason}" for rule, reason in failed_rules])
    payload = {
        "username": "Network Compliance Bot",
        "embeds": [
            {
                "title": "🚨 Security Compliance Alert: Non-Compliant Device Detected",
                "color": 15158332,  # Crimson red
                "fields": [
                    {"name": "Device Hostname", "value": device_name, "inline": True},
                    {"name": "Configuration Source", "value": source_file, "inline": True},
                    {"name": "Violated Baseline Rules", "value": failures_formatted, "inline": False}
                ],
                "footer": {"text": "Automated NetDevOps Audit Engine | Cisco Lab Audit"},
                "timestamp": datetime.utcnow().isoformat()
            }
        ]
    }

    try:
        response = requests.post(webhook_url, json=payload, timeout=5)
        if response.status_code == 204:
            print("    [✓] Security alert successfully dispatched to Discord!")
        else:
            print(f"    [-] Failed to send Discord alert. HTTP Status: {response.status_code}")
    except Exception as e:
        print(f"    [-] Error sending webhook alert: {e}")


def audit_configuration(config_text):
    """
    Analyze running-configuration text using regular expressions
    against enterprise security and baseline compliance rules.
    """
    results = {}

    # Rule 1: SSH Version 2 Check
    if re.search(r"ip\s+ssh\s+version\s+2", config_text, re.IGNORECASE):
        results["SSH_v2"] = "PASS"
    else:
        results["SSH_v2"] = "FAIL (SSH v2 Not Enforced)"

    # Rule 2: Service Password-Encryption Check
    if re.search(r"service\s+password-encryption", config_text, re.IGNORECASE):
        results["Password_Encryption"] = "PASS"
    else:
        results["Password_Encryption"] = "FAIL (Plaintext Passwords)"

    # Rule 3: Banner MOTD Check (Legal Notice)
    if re.search(r"banner\s+motd", config_text, re.IGNORECASE):
        results["Banner_MOTD"] = "PASS"
    else:
        results["Banner_MOTD"] = "FAIL (No Warning Banner)"

    # Rule 4: Telnet Disabled (Transport Input SSH only)
    if re.search(r"transport\s+input\s+ssh", config_text, re.IGNORECASE):
        results["Telnet_Disabled"] = "PASS"
    else:
        results["Telnet_Disabled"] = "FAIL (Telnet/Insecure Active)"

    return results


def generate_excel_report(audit_data, output_path):
    """
    Generate an executive-grade Excel audit report with custom styling.
    """
    df = pd.DataFrame(audit_data)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Audit_Summary")
        workbook = writer.book
        worksheet = writer.sheets["Audit_Summary"]

        pass_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
        pass_font = Font(color="155724", bold=True, name="Calibri", size=10)

        fail_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
        fail_font = Font(color="721C24", bold=True, name="Calibri", size=10)

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)

        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3")
        )
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        worksheet.row_dimensions[1].height = 28

        for row in worksheet.iter_rows(min_row=2, max_row=len(df) + 1):
            worksheet.row_dimensions[row[0].row].height = 22
            for cell in row:
                cell.border = thin_border
                cell.alignment = center_align
                val_str = str(cell.value) if cell.value is not None else ""

                if "PASS" in val_str:
                    cell.fill = pass_fill
                    cell.font = pass_font
                elif "FAIL" in val_str:
                    cell.fill = fail_fill
                    cell.font = fail_font

        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 16)

    print(f"\n[+] Styled Audit Report successfully created: {output_path}")


def main():
    print("=== Network Automated Configuration & Compliance Auditor ===")
    print("=== Mode: Cisco Packet Tracer Batch File Audit ===")

    backup_dir = "backups"
    if not os.path.exists(backup_dir):
        os.makedirs(backup_dir)

    # Get all .cfg files in backups/ folder
    cfg_files = [f for f in os.listdir(backup_dir) if f.endswith(".cfg")]
    if not cfg_files:
        print("[-] No .cfg configuration files found in 'backups/' directory!")
        return

    print(f"[+] Found {len(cfg_files)} configuration snapshot(s) to audit.")

    audit_records = []
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    for filename in cfg_files:
        filepath = os.path.join(backup_dir, filename)
        device_name = filename.replace(".cfg", "")
        print(f"\n[*] Scanning & Auditing: {device_name} ({filename})...")

        with open(filepath, "r") as f:
            config_content = f.read()

        audit_res = audit_configuration(config_content)

        # Collect failed rules
        failed_checks = []
        for rule, status in audit_res.items():
            if "FAIL" in status:
                failed_checks.append((rule, status))

        # Alert via Discord if violations found
        if failed_checks:
            print(f"    [!] Compliance violations found on {device_name} ({len(failed_checks)} issues). Dispatching alert...")
            send_discord_alert(device_name, filename, failed_checks, DISCORD_WEBHOOK_URL)
        else:
            print(f"    [✓] {device_name} is 100% compliant with security baselines.")

        record = {
            "Device Name": device_name,
            "Source File": filename,
            "Timestamp": timestamp,
            "SSH v2": audit_res["SSH_v2"],
            "Password Encryption": audit_res["Password_Encryption"],
            "Banner MOTD": audit_res["Banner_MOTD"],
            "Telnet Disabled": audit_res["Telnet_Disabled"]
        }
        audit_records.append(record)

    # Generate Excel Report
    report_file = f"reports/PT_Compliance_Audit_Report_{timestamp}.xlsx"
    generate_excel_report(audit_records, report_file)


if __name__ == "__main__":
    main()